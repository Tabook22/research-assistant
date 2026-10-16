"""Document upload validation and text extraction."""

from __future__ import annotations

import re
import uuid
from pathlib import Path

import aiofiles
from fastapi import HTTPException, UploadFile, status
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models.document import DocumentChunk, UploadedDocument
from app.models.project import ResearchProject

settings = get_settings()

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv", ".xlsx"}
CONTENT_HINTS = {
    ".pdf": ["application/pdf"],
    ".docx": [
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/octet-stream",
    ],
    ".txt": ["text/plain", "application/octet-stream"],
    ".csv": ["text/csv", "text/plain", "application/vnd.ms-excel", "application/octet-stream"],
    ".xlsx": [
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/octet-stream",
    ],
}


def _safe_name(original: str) -> str:
    base = Path(original).name
    base = re.sub(r"[^\w.\-() ]+", "_", base)
    return f"{uuid.uuid4().hex[:12]}_{base}"[:200]


async def save_upload(
    db: Session,
    project: ResearchProject,
    user_id: int,
    file: UploadFile,
    category: str = "reference_material",
) -> UploadedDocument:
    if not file.filename:
        raise HTTPException(status_code=400, detail="Filename required")

    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"File type not allowed. Allowed: {', '.join(sorted(ALLOWED_EXTENSIONS))}",
        )

    content = await file.read()
    if len(content) > settings.max_upload_bytes:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"File exceeds maximum size of {settings.max_upload_size_mb} MB",
        )
    if len(content) == 0:
        raise HTTPException(status_code=400, detail="Empty file")

    upload_root = Path(settings.upload_dir) / str(project.id)
    upload_root.mkdir(parents=True, exist_ok=True)
    safe = _safe_name(file.filename)
    path = upload_root / safe

    async with aiofiles.open(path, "wb") as f:
        await f.write(content)

    doc = UploadedDocument(
        project_id=project.id,
        uploaded_by_id=user_id,
        original_filename=file.filename,
        safe_filename=safe,
        file_type=ext.lstrip("."),
        file_size=len(content),
        storage_path=str(path),
        category=category,
        processing_status="stored",
        extracted_text_status="pending",
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    try:
        text = extract_text(path, ext)
        _store_chunks(db, doc, text)
        doc.extracted_text_status = "extracted" if text.strip() else "empty"
        doc.processing_status = "processed"
    except Exception as exc:  # noqa: BLE001
        doc.extracted_text_status = "failed"
        doc.processing_status = "error"
        doc.notes = f"Extraction error: {exc.__class__.__name__}"
    db.commit()
    db.refresh(doc)
    return doc


def extract_text(path: Path, ext: str) -> str:
    if ext == ".txt" or ext == ".csv":
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            parts.append(page.extract_text() or "")
        return "\n".join(parts)
    if ext == ".docx":
        from docx import Document

        document = Document(str(path))
        return "\n".join(p.text for p in document.paragraphs)
    if ext == ".xlsx":
        # Lightweight: store note only; full parse can be added with openpyxl
        return f"[XLSX file stored: {path.name}. Open in spreadsheet tools for full analysis.]"
    return ""


def _store_chunks(db: Session, doc: UploadedDocument, text: str, chunk_size: int = 3000) -> None:
    db.query(DocumentChunk).filter(DocumentChunk.document_id == doc.id).delete()
    if not text.strip():
        return
    for i in range(0, len(text), chunk_size):
        chunk = text[i : i + chunk_size]
        db.add(
            DocumentChunk(
                document_id=doc.id,
                chunk_index=i // chunk_size,
                content=chunk,
            )
        )
    db.commit()


def get_document_text(db: Session, doc: UploadedDocument) -> tuple[list[str], int]:
    chunks = (
        db.query(DocumentChunk)
        .filter(DocumentChunk.document_id == doc.id)
        .order_by(DocumentChunk.chunk_index)
        .all()
    )
    texts = [c.content for c in chunks]
    return texts, sum(len(t) for t in texts)
