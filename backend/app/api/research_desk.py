"""Research desk: scoped sources, grounded jobs, evidence, memory and drafts."""
from __future__ import annotations
from datetime import datetime, timezone, timedelta
from io import BytesIO
import json
import hashlib
from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException
from fastapi.responses import Response
from sqlalchemy import update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, defer
from app.database import get_db
from app.core.deps import get_current_user
from app.models.user import User
from app.models.literature import Reference
from app.models.research_desk import ResearchPassage, ResearchEvidence, ResearchBrief, ResearchJob, ResearchDraft, ResearchReferenceCheck, ResearchSourceReference
from app.schemas.research_desk import DeskTask, EvidenceCreate, EvidenceEdit, BriefEdit, DraftEdit, SourceReferenceUpdate
from app.schemas.literature import ReferenceCreate
from app.services.access import get_project_for_user
from app.services import research_retrieval as retrieval
from app.services.research_jobs import run_job
from app.services.research_discovery import verify_reference, clean_doi
from app.services.literature_service import add_reference

router=APIRouter(prefix="/projects/{project_id}/desk",tags=["research-desk"])

def owned(db,model,ident,project,user):
    get_project_for_user(db,project,user)
    row=db.get(model,ident)
    if row is None or row.project_id!=project:raise HTTPException(404,"Item not found in this project.")
    return row

def job_data(row,include_result=True):
    return {"id":row.id,"kind":row.kind,"status":row.status,"message":row.message,"payload":row.payload,"result":row.result if include_result else None,"created_at":row.created_at,"updated_at":row.updated_at}

def evidence_data(e):
    return {"id":e.id,"claim":e.claim,"quote":e.quote,"source":e.source,"notes":e.notes,"relationship":e.relationship,"version":e.version}

def draft_data(d):
    return {"id":d.id,"title":d.title,"text":d.text,"evidence":d.evidence,"version":d.version,"revisions":d.revisions,"updated_at":d.updated_at}

@router.get("/sources")
def sources(project_id:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    get_project_for_user(db,project_id,user)
    stats=retrieval.sync_index(db,project_id)
    return {"sources":retrieval.list_sources(db,project_id),"index":stats}

@router.put("/source-reference")
def link_source_reference(project_id:int,payload:SourceReferenceUpdate,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    owned(db,Reference,payload.reference_id,project_id,user)
    if not db.query(ResearchPassage).filter_by(project_id=project_id,source_key=payload.source_key).first():raise HTTPException(404,"Source not found in this project.")
    row=db.query(ResearchSourceReference).filter_by(project_id=project_id,source_key=payload.source_key).first()
    if not row:row=ResearchSourceReference(project_id=project_id,source_key=payload.source_key,reference_id=payload.reference_id);db.add(row)
    else:row.reference_id=payload.reference_id
    db.commit()
    return {"source_key":payload.source_key,"reference_id":payload.reference_id}

@router.get("/passages/{passage_id}")
def passage(project_id:int,passage_id:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    return retrieval.passage_data(owned(db,ResearchPassage,passage_id,project_id,user))

@router.post("/jobs",status_code=202)
def create_job(project_id:int,payload:DeskTask,tasks:BackgroundTasks,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    get_project_for_user(db,project_id,user)
    data=payload.model_dump()
    if payload.kind=="write":
        ids={e.id for e in db.query(ResearchEvidence).filter(ResearchEvidence.project_id==project_id,ResearchEvidence.id.in_(payload.evidence_ids)).all()}
        if ids!=set(payload.evidence_ids):raise HTTPException(404,"Evidence not found in this project.")
    for previous in db.query(ResearchJob).filter_by(project_id=project_id,kind=payload.kind).order_by(ResearchJob.id.desc()).limit(20).all():
        if previous.payload==data:
            if previous.status in ("queued","running"):return job_data(previous)
            recent=previous.created_at.replace(tzinfo=timezone.utc)>datetime.now(timezone.utc)-timedelta(hours=6)
            if payload.kind=="discover" and previous.status=="completed" and recent and all(p["status"]=="ok" for p in previous.result.get("providers",[])):
                return {**job_data(previous),"cached":True}
    if db.query(ResearchJob).filter(ResearchJob.project_id==project_id,ResearchJob.status.in_(["queued","running"])).count()>=3:raise HTTPException(429,"Three research tasks are already active. Wait or cancel one before starting another.")
    row=ResearchJob(project_id=project_id,kind=payload.kind,payload=data,status="queued",message="Queued")
    db.add(row);db.commit();db.refresh(row);tasks.add_task(run_job,row.id)
    return job_data(row)

@router.get("/jobs")
def jobs(project_id:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    get_project_for_user(db,project_id,user)
    return [job_data(r,False) for r in db.query(ResearchJob).options(defer(ResearchJob.result)).filter_by(project_id=project_id).order_by(ResearchJob.id.desc()).limit(25).all()]

@router.get("/jobs/{job_id}")
def get_job(project_id:int,job_id:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    return job_data(owned(db,ResearchJob,job_id,project_id,user))

@router.post("/jobs/{job_id}/cancel")
def cancel(project_id:int,job_id:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    row=owned(db,ResearchJob,job_id,project_id,user)
    db.execute(update(ResearchJob).where(ResearchJob.id==row.id,ResearchJob.status.in_(["running","queued"])).values(status="cancelled",message="Cancelled. An in-flight provider request may still finish; its result will not be saved."));db.commit();db.refresh(row)
    return job_data(row)

@router.get("/evidence")
def notebook(project_id:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    get_project_for_user(db,project_id,user)
    return [evidence_data(e) for e in db.query(ResearchEvidence).filter_by(project_id=project_id).order_by(ResearchEvidence.id.desc()).all()]

@router.post("/evidence",status_code=201)
def save_evidence(project_id:int,payload:EvidenceCreate,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    passage=owned(db,ResearchPassage,payload.passage_id,project_id,user)
    if passage.fingerprint!=payload.fingerprint:raise HTTPException(409,"The source changed. Search again before saving this evidence.")
    if " ".join(payload.quote.split()) not in " ".join(passage.text.split()):raise HTTPException(422,"The quotation must match the source passage.")
    if not payload.claim.strip():raise HTTPException(422,"Enter a claim or observation.")
    snapshot={**retrieval.passage_data(passage),"passage_id":passage.id}
    link=db.query(ResearchSourceReference).filter_by(project_id=project_id,source_key=passage.source_key).first()
    if link:
        ref=db.get(Reference,link.reference_id)
        if ref and ref.project_id==project_id:snapshot["reference"]={"id":ref.id,"title":ref.title,"authors":ref.authors,"year":ref.year,"doi":ref.doi,"url":ref.url,"verification_status":ref.verification_status}
    row=ResearchEvidence(project_id=project_id,claim=payload.claim,quote=payload.quote,source=snapshot,notes=payload.notes,relationship=payload.relationship)
    db.add(row);db.commit();db.refresh(row);return evidence_data(row)

@router.patch("/evidence/{evidence_id}")
def edit_evidence(project_id:int,evidence_id:int,payload:EvidenceEdit,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    row=owned(db,ResearchEvidence,evidence_id,project_id,user)
    res=db.execute(update(ResearchEvidence).where(ResearchEvidence.id==row.id,ResearchEvidence.version==payload.version).values(notes=payload.notes,relationship=payload.relationship,version=payload.version+1))
    if res.rowcount!=1:db.rollback();raise HTTPException(409,"Evidence changed in another tab. Reload it first.")
    db.commit();db.refresh(row);return evidence_data(row)

@router.get("/brief")
def brief(project_id:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    get_project_for_user(db,project_id,user);row=db.get(ResearchBrief,project_id)
    return {"content":row.content if row else "","version":row.version if row else 0}

@router.put("/brief")
def save_brief(project_id:int,payload:BriefEdit,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    get_project_for_user(db,project_id,user)
    if payload.version==0:
        try:db.add(ResearchBrief(project_id=project_id,content=payload.content));db.commit()
        except IntegrityError:db.rollback();raise HTTPException(409,"The research brief changed. Reload first.")
    else:
        res=db.execute(update(ResearchBrief).where(ResearchBrief.project_id==project_id,ResearchBrief.version==payload.version).values(content=payload.content,version=payload.version+1))
        if res.rowcount!=1:db.rollback();raise HTTPException(409,"The research brief changed. Reload first.")
        db.commit()
    row=db.get(ResearchBrief,project_id);db.refresh(row)
    return {"content":row.content,"version":row.version}

@router.post("/jobs/{job_id}/import/{item_index}")
def import_discovery(project_id:int,job_id:int,item_index:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    job=owned(db,ResearchJob,job_id,project_id,user)
    items=(job.result or {}).get("results",[])
    if job.kind!="discover" or job.status!="completed" or not 0<=item_index<len(items):raise HTTPException(404,"Discovery result not found.")
    item=items[item_index];doi=clean_doi(item.get("doi"))
    for ref in db.query(Reference).filter_by(project_id=project_id).all():
        if (doi and clean_doi(ref.doi)==doi) or ref.title.casefold()==item["title"].casefold():return {"id":ref.id,"already_saved":True}
    project=get_project_for_user(db,project_id,user)
    ref=add_reference(db,project,ReferenceCreate(title=item["title"],authors=item.get("authors"),year=item.get("year"),journal=item.get("journal"),doi=doi or None,url=item.get("url"),verification_status="metadata_retrieved",notes="Discovered via "+", ".join(item["providers"])+". Full text has not been reviewed."))
    return {"id":ref.id,"already_saved":False}

@router.post("/references/{reference_id}/verify")
async def verify(project_id:int,reference_id:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    row=owned(db,Reference,reference_id,project_id,user)
    fingerprint=hashlib.sha256((row.title+"\0"+(row.doi or "")).encode()).hexdigest()
    recent=db.query(ResearchReferenceCheck).filter_by(project_id=project_id,reference_id=row.id,fingerprint=fingerprint).order_by(ResearchReferenceCheck.id.desc()).first()
    if recent and recent.created_at.replace(tzinfo=timezone.utc)>datetime.now(timezone.utc)-timedelta(hours=24) and recent.result.get("status")!="provider_unavailable":
        return {**recent.result,"cached":True}
    result={**await verify_reference(row),"checked_at":datetime.now(timezone.utc).isoformat(),"reference_id":row.id}
    db.add(ResearchReferenceCheck(project_id=project_id,reference_id=row.id,fingerprint=fingerprint,result=result));db.commit()
    return result

@router.get("/reference-checks")
def saved_reference_checks(project_id:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    get_project_for_user(db,project_id,user)
    refs={r.id:r for r in db.query(Reference).filter_by(project_id=project_id).all()}
    checks={}
    for row in db.query(ResearchReferenceCheck).filter_by(project_id=project_id).order_by(ResearchReferenceCheck.id.desc()).all():
        ref=refs.get(row.reference_id)
        if not ref or row.reference_id in checks:continue
        fingerprint=hashlib.sha256((ref.title+"\0"+(ref.doi or "")).encode()).hexdigest()
        if fingerprint==row.fingerprint:checks[row.reference_id]=row.result
    return checks

@router.post("/jobs/{job_id}/draft",status_code=201)
def keep_draft(project_id:int,job_id:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    job=owned(db,ResearchJob,job_id,project_id,user)
    if job.kind!="write" or job.status!="completed" or not (job.result or {}).get("claims"):raise HTTPException(422,"No supported writing proposal is available.")
    text="\n\n".join(c["text"]+" "+" ".join(f"[E{x['id']}]" for x in c["citations"]) for c in job.result["claims"])
    row=ResearchDraft(project_id=project_id,title=job.payload["query"][:255],text=text,evidence=job.result["evidence_snapshot"],revisions=[])
    db.add(row);db.commit();db.refresh(row);return draft_data(row)

@router.get("/drafts")
def drafts(project_id:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    get_project_for_user(db,project_id,user)
    return [draft_data(d) for d in db.query(ResearchDraft).filter_by(project_id=project_id).order_by(ResearchDraft.id.desc()).all()]

@router.patch("/drafts/{draft_id}")
def edit_draft(project_id:int,draft_id:int,payload:DraftEdit,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    row=owned(db,ResearchDraft,draft_id,project_id,user)
    revisions=[*row.revisions,{"text":row.text,"title":row.title,"saved_at":row.updated_at.isoformat()}]
    res=db.execute(update(ResearchDraft).where(ResearchDraft.id==row.id,ResearchDraft.version==payload.version).values(text=payload.text,title=payload.title,revisions=revisions,version=payload.version+1))
    if res.rowcount!=1:db.rollback();raise HTTPException(409,"Draft changed in another tab. Reload before saving.")
    db.commit();db.refresh(row);return draft_data(row)

@router.get("/drafts/{draft_id}/export")
def export_draft(project_id:int,draft_id:int,format:str="markdown",db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    row=owned(db,ResearchDraft,draft_id,project_id,user)
    appendix=[]
    for e in row.evidence:
        s=e["source"];ref=s.get("reference") or {}
        citation=f"{ref.get('authors') or ''} ({ref.get('year') or 'n.d.'}). {ref.get('title') or s['title']}. {('https://doi.org/'+ref['doi']) if ref.get('doi') else ref.get('url') or ''}"
        appendix.append(f"[E{e['id']}] {citation} — {s['location']}\nSource excerpt: {e['quote']}\nResearcher notes: {e['notes']}")
    if format=="docx":
        from docx import Document
        doc=Document();doc.add_heading(row.title,0)
        for paragraph in row.text.split("\n\n"):doc.add_paragraph(paragraph)
        doc.add_heading("Evidence and source locations",1)
        doc.add_paragraph("Evidence markers identify saved excerpts; they are not a formatted scholarly bibliography.")
        for text in appendix:doc.add_paragraph(text)
        stream=BytesIO();doc.save(stream)
        return Response(stream.getvalue(),media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",headers={"Content-Disposition":f'attachment; filename="research-draft-{row.id}.docx"'})
    if format!="markdown":raise HTTPException(422,"Choose markdown or docx.")
    content="# "+row.title+"\n\n"+row.text+"\n\n## Evidence and source locations\n\n"+"\n\n".join(appendix)
    return Response(content,media_type="text/markdown; charset=utf-8",headers={"Content-Disposition":f'attachment; filename="research-draft-{row.id}.md"'})


@router.get("/references/export")
def export_reference_records(project_id:int,db:Session=Depends(get_db),user:User=Depends(get_current_user)):
    get_project_for_user(db,project_id,user)
    rows=db.query(Reference).filter_by(project_id=project_id).all()
    clean=lambda value:" ".join(str(value or "").splitlines())
    records=[]
    for ref in rows:
        lines=["TY  - "+("JOUR" if ref.journal else "GEN"),"TI  - "+clean(ref.title)]
        if ref.year:lines.append("PY  - "+str(ref.year))
        if ref.doi:lines.append("DO  - "+clean(clean_doi(ref.doi)))
        if ref.url:lines.append("UR  - "+clean(ref.url))
        if ref.journal:lines.append("JO  - "+clean(ref.journal))
        fingerprint=hashlib.sha256((ref.title+"\0"+(ref.doi or "")).encode()).hexdigest()
        check=db.query(ResearchReferenceCheck).filter_by(project_id=project_id,reference_id=ref.id,fingerprint=fingerprint).order_by(ResearchReferenceCheck.id.desc()).first()
        authors=(check.result.get("authors") or []) if check and check.result.get("status")=="metadata_found" else []
        for author in authors:
            name=(author.get("family","")+", "+author.get("given","")).strip(", ") or author.get("name","")
            if name:lines.append("AU  - "+clean(name))
        if not authors and ref.authors:lines.append("N1  - Author metadata needs checking. Original author text: "+clean(ref.authors))
        lines.append("N1  - Exported metadata; confirm bibliographic fields before citing.")
        lines.append("ER  - ");records.append("\n".join(lines))
    return Response("\n\n".join(records),media_type="application/x-research-info-systems; charset=utf-8",headers={"Content-Disposition":'attachment; filename="research-references.ris"'})
